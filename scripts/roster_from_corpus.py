"""从范例语料里读「班级名册」—— 每个幼儿的康复目标 / 康复基线 / 参与情况。

## 为什么需要这个模块

实测 14 份范例、189 条儿童记录：

1. `康复基线` 列**空值 0 条**；
2. **记号个数 == 该幼儿「康复目标」的编号个数，178/178 无例外**
   （目标 `1456` 配 `P  P PP`，目标 `145` 配 `P  EE`）；
3. 同一幼儿在连续教案里是**同一串**，说明它是**幼儿的基线能力描述**、
   随教案往下传递。

所以新教案里每位幼儿的这两个值，应当**取自他/她最近一次有记录的教案**。

## 「未参与」的判定（用户要求：统一标注为 `未参与`）

范例里 `康复基线` 会出现 `未参与`，实测触发条件 = **该幼儿整轮未参与**：

    第三周-浓浓端午节  蒙智宇  基线 未参与   五次评量 全 未参与
    第四周-清凉夏日    蒙智宇  基线 未参与   五次评量 全 未参与
    第五周-童心向党    蒙智宇  基线 PEEE     五次评量 全 未参与   ← 老师漏改

反例（只缺某几次，基线照常）：何侦霖 `P  EE`、陈映彤 `P  EE`、覃芯玥 `PEE`。

所以本模块的规则：
- 该幼儿**最近一次记录**的五次评量**全为「未参与」**（或基线格直接写「未参与」）
  → 基线 = `未参与`；
- 否则 → 用记录里的基线串原样。
被 `attendance` 判为「缺席」的幼儿，正是「所有记录都整轮未参与」的那几位。

## 行定位

主表 = 行数最多的那张表。数据行 = 「儿童姓名」列（tc 序号 3）非空**且命中真名白名单**。
⚠ 不能只看「姓名非空」—— 表头行 `儿童姓名`、统计行 `人数`、`13` 等都会被误认。

⚠ 也不能按「康复目标必须是纯数字」筛行 —— 《第四周-能干的我》《第四周-浓浓端午节》
里有 4 行的康复目标格被 P/E 串挤串（读到 `P   PP`），按数字筛会整行丢掉。
"""
from __future__ import annotations

import json
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
CORPUS = Path(r"C:\Users\17093\Desktop\homework\agent use idea")

# 主表里各逻辑列的 tc 序号（0 起）
COL_NAME = 3      # 儿童姓名
COL_GOAL = 4      # 康复目标（编号串，如 1456）
COL_BASE = 5      # 康复基线（P/E 串，如 P  P PP）
COL_EVAL1 = 6     # 第一次评量 … 第五次 = 6..10
N_EVAL = 5

ABSENT_MARK = "未参与"


def known_names() -> set:
    """姓名白名单 —— 私有映射里的 37 个真名（仅内部用，绝不写进产物）。"""
    p = ROOT / "data" / "private" / "name_mapping.json"
    if not p.exists():
        return set()
    return set(json.loads(p.read_text(encoding="utf-8")))


def parse_start_date(date_range: str) -> tuple:
    """把 `2025年5月6日—5月9日` / `2026年6月29日--7月3日` 解析成 (y, m, d)。"""
    if not date_range:
        return (0, 0, 0)
    m = re.search(r"(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日", date_range)
    if m:
        return (int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = re.search(r"(\d{4})[-/](\d{1,2})[-/](\d{1,2})", date_range)
    if m:
        return (int(m.group(1)), int(m.group(2)), int(m.group(3)))
    return (0, 0, 0)


def _cell_text(tc) -> str:
    from docx.oxml.ns import qn
    return "".join(n.text or "" for n in tc.iter(qn("w:t"))).strip()


def read_lesson_rows(fp: Path, whitelist: set = None) -> list:
    """读一份范例的儿童数据行。

    → [{name, goal, base, evals, all_absent}, ...]
    """
    from docx import Document
    from docx.oxml.ns import qn

    try:
        d = Document(str(fp))
    except Exception:
        return []
    if not d.tables:
        return []
    t = max(d.tables, key=lambda tb: len(tb.rows))
    out: list = []
    for row in t.rows:
        tcs = row._tr.findall(qn("w:tc"))
        if len(tcs) <= COL_EVAL1 + N_EVAL - 1:
            continue
        nm = _cell_text(tcs[COL_NAME])
        if not nm or (whitelist and nm not in whitelist):
            continue
        goal = _cell_text(tcs[COL_GOAL])
        base = _cell_text(tcs[COL_BASE])
        # 被挤串时：goal 格会读到 P/E 串、base 格为空 → 回退
        if goal and not any(ch.isdigit() for ch in goal):
            base = base or goal
            goal = ""
        evals = [_cell_text(tcs[COL_EVAL1 + k]) for k in range(N_EVAL)]
        nonempty = [e for e in evals if e]
        all_absent = bool(nonempty) and all(ABSENT_MARK in e for e in nonempty)
        out.append({
            "name": nm, "goal": goal, "base": base,
            "evals": evals, "all_absent": all_absent,
        })
    return out


def load_lessons(whitelist: set = None) -> list:
    """→ [{file, title, date_range, start, rows}, ...]，按日期升序。"""
    meta = json.loads(
        (ROOT / "data" / "structured" / "lessons.json").read_text(encoding="utf-8"))
    out: list = []
    for L in meta:
        fp = CORPUS / L["source_file"]
        if not fp.exists() or fp.name.startswith("~$"):
            continue
        rows = read_lesson_rows(fp, whitelist)
        out.append({
            "file": L["source_file"],
            "title": L.get("title", ""),
            "date_range": L.get("date_range", ""),
            "start": parse_start_date(L.get("date_range", "")),
            "rows": rows,
        })
    out.sort(key=lambda x: (x["start"], x["file"]))
    return out


def normalize_baseline(rec: dict) -> str:
    """把一条记录的「康复基线」规整成可写入教案的值。

    规则（用户要求「未参与的情况统一标注为未参与」）：
    - 整轮未参与 → `未参与`；
    - 基线格本身写着「未参与」→ `未参与`；
    - 有 P/E 记号 → 原样保留（含原表的对位空格）；
    - 既无记号也不是未参与（空值）→ `未参与`（无有效基线可比）。
    """
    if not rec:
        return ABSENT_MARK
    if rec.get("all_absent"):
        return ABSENT_MARK
    raw = (rec.get("base") or "").strip()
    if ABSENT_MARK in raw:
        return ABSENT_MARK
    if not any(ch in "PE" for ch in raw):
        return ABSENT_MARK
    return raw


def build_history(whitelist: set = None) -> dict:
    """→ {姓名: [记录, ...]}，每人的出现序列（按日期升序）。"""
    wl = whitelist if whitelist is not None else known_names()
    hist: dict = {}
    for L in load_lessons(wl):
        for rec in L["rows"]:
            item = dict(rec)
            item.update({"file": L["file"], "title": L["title"],
                         "date_range": L["date_range"], "start": L["start"]})
            hist.setdefault(rec["name"], []).append(item)
    return hist


def latest_for(hist: dict, name: str) -> dict:
    """该幼儿**最新日期**那次记录（即「包含其姓名的最新日期范例」）。"""
    seq = hist.get(name) or []
    return seq[-1] if seq else {}


def roster_records(anon2name: dict, hist: dict = None) -> dict:
    """→ {anon_id: {name, goal, baseline, date_range, source, absent}}。

    `anon2name` 是 `Sxx → 真名`（私有，仅内部取数）。
    """
    hist = hist if hist is not None else build_history()
    out: dict = {}
    for aid, nm in anon2name.items():
        rec = latest_for(hist, nm)
        out[aid] = {
            "name": nm,
            "goal": (rec.get("goal") or "").strip(),
            "baseline": normalize_baseline(rec),
            "date_range": rec.get("date_range", ""),
            "source": rec.get("file", ""),
            "absent": normalize_baseline(rec) == ABSENT_MARK,
        }
    return out


if __name__ == "__main__":
    anon2name = {}
    m = json.loads((ROOT / "data" / "private" / "name_mapping.json")
                   .read_text(encoding="utf-8"))
    anon2name = {v: k for k, v in m.items()}
    recs = roster_records(anon2name)
    print(f"名册 {len(recs)} 人\n")
    for aid in sorted(recs):
        r = recs[aid]
        print(f"{aid} {r['name']:5s} 目标={r['goal']:5s} 基线={r['baseline']!r:11s} "
              f"{'[未参与]' if r['absent'] else '        '} ← {r['date_range']}")
