# -*- coding: utf-8 -*-
"""抽取教师原表格的**完整模板结构**，作为 W3 docx 渲染的还原目标。

输出两个东西：
  1. 每份 docx 的两张表：表0（活动名称）+ 表1（主表）的**逐行逐格结构**
  2. 一份「主流模板骨架」—— 各份表1的共有行标签序列

不改任何数据，只做结构 dump。
"""

from __future__ import annotations

import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from lpa.template import tc_cells_raw, table_shape  # noqa: E402

SRC = Path(r"C:\Users\17093\Desktop\homework\agent use idea")
ROOT = Path(r"C:\Users\17093\WorkBuddy\2026-09-28-10-57-57\lessonplan-agent")
OUT = ROOT / "data" / "structured" / "TEMPLATE_STRUCTURE.md"

files = sorted(f for f in SRC.iterdir()
               if f.suffix.lower() == ".docx" and not f.name.startswith("~$"))

lines = ["# 教师原表格模板结构（逐份 dump）", ""]
row_label_seqs = Counter()

for f in files:
    doc = Document(str(f))
    lines += [f"## {f.name}", ""]
    for ti, t in enumerate(doc.tables):
        shape = table_shape(t)
        lines.append(f"**表{ti}**：{shape['n_rows']} 行 × {shape['n_cols']} 列（按 gridSpan 展开后）")
        lines.append("")
        lines.append("| 行 | " + " | ".join(f"c{i}" for i in range(min(shape["n_cols"], 12))) + " |")
        lines.append("|" + "---|" * (min(shape["n_cols"], 12) + 1))
        for ri, row in enumerate(t.rows):
            cells = tc_cells_raw(row)
            disp = []
            for c in cells[:12]:
                c = c.replace("\n", "⏎")
                disp.append((c[:16] + "…") if len(c) > 16 else (c or "·"))
            while len(disp) < min(shape["n_cols"], 12):
                disp.append("·")
            lines.append(f"| {ri} | " + " | ".join(disp) + " |")
        lines.append("")
    # 记录表1的行标签序列（第一个非空格）
    if len(doc.tables) >= 2:
        labels = []
        for row in doc.tables[1].rows:
            cells = tc_cells_raw(row)
            for c in cells:
                fc = c.replace("\n", "").strip()
                if fc and len(fc) <= 8:
                    labels.append(fc)
                    break
        row_label_seqs[tuple(labels)] += 1

lines += ["", "## 主流表1行标签序列", ""]
for seq, n in row_label_seqs.most_common():
    lines += [f"**出现 {n} 次**：", "", "```", " / ".join(seq), "```", ""]

OUT.write_text("\n".join(lines), encoding="utf-8")
print("ok", len(files), "files")
