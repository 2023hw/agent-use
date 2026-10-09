"""检查 W3 产物：把 docx 的可读文本 dump 到控制台，便于人工核对。"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from docx import Document                       # noqa: E402
from lpa.render_docx import tc_cells_raw        # noqa: E402


def main() -> None:
    files = sorted((ROOT / "data" / "docx").glob("*.docx"))
    if not files:
        print("no docx")
        return
    f = files[0]
    d = Document(str(f))
    print(f"文件：{f.name}")
    print("表外段落：", [p.text for p in d.paragraphs if p.text.strip()])
    print("表数：", len(d.tables))
    for ti, t in enumerate(d.tables):
        print(f"\n--- 表{ti}（{len(t.rows)} 行）---")
        for ri, row in enumerate(t.rows):
            raw = [x for x in tc_cells_raw(row) if x]
            print(f"[r{ri}]")
            for i, x in enumerate(tc_cells_raw(row)):
                if x:
                    print(f"   tc{i}: {x[:200]}")


if __name__ == "__main__":
    main()
