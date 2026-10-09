# -*- coding: utf-8 -*-
"""从原始 docx 教案中提取纯文本，落盘为 JSON，便于后续分析字段结构。"""
import json
import sys
from pathlib import Path

from docx import Document

SRC = Path(r"C:\Users\17093\Desktop\homework\agent use idea")
OUT = Path(r"C:\Users\17093\WorkBuddy\2026-09-28-10-57-57\lessonplan-agent\data\extracted")
OUT.mkdir(parents=True, exist_ok=True)


def extract_docx(path: Path):
    doc = Document(str(path))
    blocks = []
    # 正文段落
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            blocks.append({"type": "para", "text": t})
    # 表格
    for ti, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            rows.append(cells)
        blocks.append({"type": "table", "index": ti, "rows": rows})
    return blocks


def main():
    result = []
    for f in sorted(SRC.iterdir()):
        if f.suffix.lower() != ".docx":
            continue
        try:
            blocks = extract_docx(f)
        except Exception as e:
            blocks = [{"type": "error", "text": f"{type(e).__name__}: {e}"}]
        result.append({
            "file": f.name,
            "n_blocks": len(blocks),
            "blocks": blocks,
        })
        print(f"{f.name}: {len(blocks)} blocks", flush=True)
    out_file = OUT / "raw_text.json"
    out_file.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print("WROTE", out_file)


if __name__ == "__main__":
    main()
