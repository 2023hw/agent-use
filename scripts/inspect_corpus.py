# -*- coding: utf-8 -*-
"""摸清语料库骨架：每份教案有哪些字段、是否含个别化目标表、是否有儿童姓名。"""
import json
import re
from pathlib import Path

from docx import Document

SRC = Path(r"C:\Users\17093\Desktop\homework\agent use idea")
OUT = Path(r"C:\Users\17093\WorkBuddy\2026-09-28-10-57-57\lessonplan-agent\data\extracted")

NAME_HINT = re.compile(r"^[\u4e00-\u9fa5]{2,4}$")


def uniq_cells(rows):
    """返回每行第一列的去重短文本，用于识别字段名。"""
    labels = []
    for row in rows:
        if not row:
            continue
        first = row[0].strip()
        if not first:
            continue
        first = re.sub(r"\s+", " ", first)[:24]
        if first not in labels:
            labels.append(first)
    return labels


def scan(path: Path):
    doc = Document(str(path))
    info = {
        "file": path.name,
        "paragraphs": len([p for p in doc.paragraphs if p.text.strip()]),
        "tables": len(doc.tables),
        "first_paragraph": next((p.text.strip() for p in doc.paragraphs if p.text.strip()), ""),
    }
    headers = []
    for t in doc.tables:
        headers.extend(uniq_cells([[c.text.strip() for c in r.cells] for r in t.rows]))
    # 去重保序
    seen, uniq = set(), []
    for h in headers:
        if h not in seen:
            seen.add(h)
            uniq.append(h)
    info["field_labels"] = uniq

    # 提取形如「儿童姓名」表里的姓名 + 评量码
    names, codes = [], []
    for t in doc.tables:
        for r in t.rows:
            cells = [c.text.strip() for c in r.cells]
            joined = "".join(cells)
            m = re.findall(r"\b[P E]{1,6}\b", joined.replace(" ", ""))
            if re.fullmatch(r"[PE]{2,8}", joined.replace(" ", "")):
                codes.append(joined.replace(" ", ""))
            for c in cells:
                if NAME_HINT.match(c) and c not in ("活动班级", "有益经验", "材料准备", "教学过程"):
                    names.append(c)
    # 姓名去重保序
    seen = set()
    names = [n for n in names if not (n in seen or seen.add(n))]
    info["sample_names"] = names[:24]
    info["n_evaluation_codes"] = len(codes)
    info["evaluation_code_samples"] = codes[:6]

    full = json.dumps(info, ensure_ascii=False)
    info["has_iep_table"] = "个别化" in full
    return info


def main():
    out = []
    for f in sorted(SRC.iterdir()):
        if f.suffix.lower() != ".docx":
            continue
        try:
            out.append(scan(f))
        except Exception as e:
            out.append({"file": f.name, "error": f"{type(e).__name__}: {e}"})
    target = OUT / "corpus_overview.json"
    target.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")

    for item in out:
        print("=" * 70)
        print(item.get("file"))
        if "error" in item:
            print("  ERROR:", item["error"])
            continue
        print("  段落数:", item["paragraphs"], "| 表格数:", item["tables"],
              "| 个别化表:", item["has_iep_table"],
              "| 评量码行:", item["n_evaluation_codes"])
        print("  首段:", item["first_paragraph"][:60])
        print("  字段:", " / ".join(item["field_labels"][:14]))

    print("\nWROTE", target)


if __name__ == "__main__":
    main()
