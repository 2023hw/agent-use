# -*- coding: utf-8 -*-
"""M4-5 教师盲评 · 可填写 Word 评分卡（生成 + 回收解析）。

为什么要有这一版
----------------
早先的材料包是「Markdown + 打印版 HTML」：教师只能在纸上勾选，
组织方再手工把分数誊进 CSV —— 两道人工转录，既慢又容易抄错。
这里把评分卡做成**可直接在电脑上填**的 .docx：

* 打分格是**下拉选择**（1 / 1.5 / … / 5），点一下就能选，不用手打；
* Q6 是下拉四选一；Q7 / Q8 直接在格子里打字；
* 每份教案的**正文就在该份评分表的正上方**，读到哪儿评到哪儿；
* 教师填完把文件丢进 `回收/` 文件夹，`--ingest` 自动读回并落进
  `3_回收统计表.csv`，**不需要任何人工转录**。

设计取舍
--------
下拉控件用的是 OOXML 内容控件（`w:sdt` + `w:dropDownList`）：
Word 2013+ 与较新版 WPS 都支持「点击展开列表」。**当前选中的值以纯文本
存在 `sdtContent` 里**，所以解析时按文本读即可 —— 不依赖控件本身是否被识别。
万一教师的软件不认控件，格子会退化成一段普通文字（默认「—」/「请选择」），
直接改写就行；解析逻辑同样认「手打的 4」「手打的 4.5」。

⚠ 生成的材料同样受项目合规红线约束：零幼儿姓名、零诊断标签、零来源指纹。
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Optional, Sequence

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

# --------------------------------------------------------------------------
# 常量
# --------------------------------------------------------------------------
LATIN_FONT = "Times New Roman"
CJK_FONT = "宋体"

FORM_FONT_PT = 10.5          # 五号，表单常用字号
TITLE_FONT_PT = 15.0

DIMS = [
    ("D1", "目标可测性", "活动目标是否具体、可观察、可评价；是否与主题、与幼儿年龄相称。"),
    ("D2", "分层与个别化", "是否照顾到不同能力水平的幼儿；支持方式是否由「多辅助」到"
                          "「少辅助」形成梯度；个别化目标是否落得下去。"),
    ("D3", "环节完整性", "教学过程是否完整（引入—展开—结束）、层次清楚；各环节时间分配是否合理。"),
    ("D4", "材料与环境", "材料准备是否具体、拿得到、用得上；环境布置与安全提示是否考虑到位。"),
    ("D5", "教师语言与可操作性", "教师语言能否照着说；指令是否清楚、适合幼儿；拿到手就能上课的程度。"),
]
DIM_CODES = [d[0] for d in DIMS]

# 下拉选项：1–5 分，允许半分
SCORE_OPTIONS = ["—"] + [f"{x / 2:g}" for x in range(2, 11)]   # —,1,1.5,…,5

BLANK = {"", "—", "-", "请选择", "（请在此输入）"}

Q6_OPTIONS = ["请选择", "可以直接使用", "小改一下就能用", "需要较大修改", "基本不能用"]

ROW_LABELS = {
    "score": "得分（1–5）",
    "q6": "Q6 能否直接用",
    "q7": "Q7 最需改进",
    "q8": "Q8 最有价值",
}
TEACHER_LABELS = {"teacher": "教师称呼（如：张老师）",
                  "exp": "教龄 / 专业背景",
                  "mixed": "是否带过混合能力班级"}

PLACEHOLDER_ANSWER = "（请在此输入）"


# --------------------------------------------------------------------------
# 基础排版工具
# --------------------------------------------------------------------------
def _font(run, size_pt: Optional[float] = None, bold: Optional[bool] = None) -> None:
    """把 run 刷成「Times New Roman + 宋体」。

    ⚠ 不要用 `run.font.name`：它只写 `w:ascii`/`w:hAnsi`，会把西文也设成中文字体名。
      显式写全四个属性 + hint，中英各归其位。
    """
    r = getattr(run, "_element", run)
    rpr = r.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    rf.set(qn("w:hint"), "eastAsia")
    rf.set(qn("w:ascii"), LATIN_FONT)
    rf.set(qn("w:hAnsi"), LATIN_FONT)
    rf.set(qn("w:eastAsia"), CJK_FONT)
    rf.set(qn("w:cs"), LATIN_FONT)
    if size_pt is not None:
        half = str(int(round(size_pt * 2)))
        rpr.get_or_add_sz().set(qn("w:val"), half)
        if hasattr(rpr, "get_or_add_szCs"):
            rpr.get_or_add_szCs().set(qn("w:val"), half)
    if bold is not None:
        run.font.bold = bold


def _init_styles(doc) -> None:
    """把 Normal 样式刷成表单字体档（新文档没有可继承的模板父本）。"""
    try:
        st = doc.styles["Normal"]
    except KeyError:                                       # pragma: no cover
        return
    rpr = st.element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    rf.set(qn("w:hint"), "eastAsia")
    rf.set(qn("w:ascii"), LATIN_FONT)
    rf.set(qn("w:hAnsi"), LATIN_FONT)
    rf.set(qn("w:eastAsia"), CJK_FONT)
    rf.set(qn("w:cs"), LATIN_FONT)
    rpr.get_or_add_sz().set(qn("w:val"), str(int(FORM_FONT_PT * 2)))
    st.paragraph_format.space_after = Pt(4)


def _para(doc_or_cell, text: str = "", size_pt: Optional[float] = None,
          bold: bool = False, align: Optional[int] = None):
    p = doc_or_cell.add_paragraph()
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        _font(r, size_pt, bold)
    return p


def _rich(p, text: str, size_pt: Optional[float] = None, bold: bool = False) -> None:
    """把带 `**` 的文本写成真加粗的多 run。"""
    for seg in re.split(r"(\*\*.+?\*\*)", text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**") and len(seg) > 4:
            r = p.add_run(seg[2:-2])
            _font(r, size_pt, True)
        else:
            r = p.add_run(seg)
            _font(r, size_pt, bold)


def _set_cell_lines(cell, lines: Sequence[str], size_pt: Optional[float] = None,
                    bold: bool = False) -> None:
    """清空单元格，写入多行文本（每个元素一行，行内支持 `**`）。"""
    cell.text = ""
    first = cell.paragraphs[0]
    for i, line in enumerate(lines):
        p = first if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        _rich(p, str(line), size_pt, bold)


def _row_height(row, twips: int) -> None:
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(twips))
    h.set(qn("w:hRule"), "atLeast")
    trPr.append(h)


def _make_table(doc, rows: int, cols: int):
    tbl = doc.add_table(rows=rows, cols=cols)
    try:
        tbl.style = doc.styles["Table Grid"]
    except KeyError:                                       # pragma: no cover
        pass
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    return tbl


# --------------------------------------------------------------------------
# 内容控件：下拉列表（w:sdt + w:dropDownList）
# --------------------------------------------------------------------------
def _dropdown(alias: str, options: Sequence[str], default: int = 0):
    """造一个下拉内容控件元素。当前值以纯文本落在 `sdtContent` 里 —— 可解析。

    ⚠ `w:sdtPr` 里 `w:rPr` 必须是**第一个**子元素（CT_SdtPr 是顺序约束），
      写到后面 Word 会报「文件已损坏」。
    """
    items = "".join(
        f'<w:listItem w:displayText="{o}" w:value="{o}"/>' for o in options
    )
    half = str(int(FORM_FONT_PT * 2))
    rpr = (
        f"<w:rPr>"
        f'<w:rFonts w:hint="eastAsia" w:ascii="{LATIN_FONT}" w:hAnsi="{LATIN_FONT}"'
        f' w:eastAsia="{CJK_FONT}" w:cs="{LATIN_FONT}"/>'
        f'<w:sz w:val="{half}"/><w:szCs w:val="{half}"/>'
        f"</w:rPr>"
    )
    xml = (
        f'<w:sdt {nsdecls("w")}>'
        f"<w:sdtPr>{rpr}"
        f'<w:alias w:val="{alias}"/>'
        f'<w:tag w:val="{alias}"/>'
        f"<w:dropDownList>{items}</w:dropDownList>"
        f"</w:sdtPr>"
        f"<w:sdtContent>"
        f'<w:r><w:t xml:space="preserve">{options[default]}</w:t></w:r>'
        f"</w:sdtContent>"
        f"</w:sdt>"
    )
    return parse_xml(xml)


def _put_dropdown(cell, alias: str, options: Sequence[str], default: int = 0) -> None:
    """把下拉控件放进单元格（保留一个空段落承载）。"""
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    # 先写一个空的 rPr 承载体，让控件继承表单字体档
    holder = p.add_run("")
    _font(holder, FORM_FONT_PT)
    p._p.append(_dropdown(alias, options, default))


def _deep_text(el) -> str:
    """取元素下**所有** `w:t` 文本（含内容控件内部）。

    python-docx 的 `cell.text` / `paragraph.text` 只认 `w:p` 的直接子 run，
    内容控件里的 run 在 `w:sdtContent` 包裹下，会被整段漏掉 —— 必须深挖。
    """
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


# --------------------------------------------------------------------------
# Markdown 卡片 → docx
# --------------------------------------------------------------------------
def _md_blocks(md: str):
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            i += 1
            continue
        if ln.lstrip().startswith("|"):
            tbl = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                tbl.append(lines[i])
                i += 1
            yield ("table", tbl)
        elif re.match(r"^#{1,4} ", ln):
            depth = len(ln) - len(ln.lstrip("#"))
            yield ("h", (min(depth, 4), ln.lstrip("#").strip()))
            i += 1
        elif ln.lstrip().startswith("> "):
            yield ("q", ln.strip()[2:].strip())
            i += 1
        elif re.match(r"^\s*[-*] ", ln):
            yield ("li", re.sub(r"^\s*[-*] ", "", ln).strip())
            i += 1
        else:
            yield ("p", ln.strip())
            i += 1


_HEAD_PT = {1: 14.0, 2: 12.0, 3: 11.0, 4: FORM_FONT_PT}


def _md_into(doc, md: str) -> None:
    """把一张教案卡片（Markdown）渲染进 docx 正文。"""
    for kind, payload in _md_blocks(md):
        if kind == "h":
            depth, text = payload
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8 if depth <= 2 else 5)
            p.paragraph_format.space_after = Pt(3)
            _rich(p, text, _HEAD_PT.get(depth, FORM_FONT_PT), bold=True)
        elif kind == "q":
            p = doc.add_paragraph()
            r = p.add_run(payload)
            _font(r, FORM_FONT_PT)
            r.font.italic = True
        elif kind == "li":
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(1)
            _rich(p, payload, FORM_FONT_PT)
        elif kind == "p":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            _rich(p, payload, FORM_FONT_PT)
        elif kind == "table":
            rows = [r for r in payload if not re.match(r"^\|[\s:\-|]+\|$", r.strip())]
            if not rows:
                continue
            grid = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
            ncol = max(len(r) for r in grid)
            tbl = _make_table(doc, len(grid), ncol)
            for ri, cells in enumerate(grid):
                for ci in range(ncol):
                    txt = cells[ci] if ci < len(cells) else ""
                    _set_cell_lines(tbl.rows[ri].cells[ci], [txt] if txt else [""],
                                    FORM_FONT_PT, bold=(ri == 0))
            tbl.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            _para(doc, "", FORM_FONT_PT)


# --------------------------------------------------------------------------
# 一页式填写说明
# --------------------------------------------------------------------------
def _cover(doc, codes: Sequence[str], teacher_label: str) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("教案评测量表（教师用表）")
    _font(r, TITLE_FONT_PT, True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = f"共 {len(codes)} 份（教案 {codes[0]}–{codes[-1]}）"
    if teacher_label and teacher_label not in ("空白模板", "模板"):
        sub = f"本表编号：{teacher_label}　|　" + sub
    r = p.add_run(sub)
    _font(r, FORM_FONT_PT)

    _para(doc, "", FORM_FONT_PT)
    p = doc.add_paragraph()
    _rich(p, "**怎么填（本文件可直接用电脑填）**", FORM_FONT_PT)
    for line in [
        "1. 每份教案的正文就在**它自己那张评分表的正上方**，读完随手打分，不用来回翻页。",
        "2. 带**下拉箭头**的格子，用鼠标点一下会出现选项，选一个即可；"
        "若您的软件点不开，直接删掉格子里的字、手打也可以（例如 4 或 4.5）。",
        "3. **D1–D5** 每项打 1–5 分，允许半分；**Q6** 选一项；**Q7** 必答、**Q8** 选答，"
        "在格子里输入文字。",
        "4. 全程**匿名、自愿**，您可以随时停止，不需要理由。",
        "5. 填完后保存本文件，改名为「**评分卡-您的称呼.docx**」，"
        "放进**回收**文件夹交给组织方即可。",
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(2)
        _rich(p, line, FORM_FONT_PT)

    _para(doc, "", FORM_FONT_PT)
    p = doc.add_paragraph()
    _rich(p, "**评分维度**", FORM_FONT_PT)
    tbl = _make_table(doc, len(DIMS) + 1, 3)
    header = ["编号", "维度", "看什么"]
    for ci, h in enumerate(header):
        _set_cell_lines(tbl.rows[0].cells[ci], [h], FORM_FONT_PT, bold=True)
    for ri, (code, name, desc) in enumerate(DIMS, start=1):
        _set_cell_lines(tbl.rows[ri].cells[0], [code], FORM_FONT_PT, bold=True)
        _set_cell_lines(tbl.rows[ri].cells[1], [name], FORM_FONT_PT)
        _set_cell_lines(tbl.rows[ri].cells[2], [desc], FORM_FONT_PT)
    for ci, w in enumerate((Cm(1.4), Cm(3.0), Cm(11.6))):
        for row in tbl.rows:
            row.cells[ci].width = w

    _para(doc, "", FORM_FONT_PT)
    p = doc.add_paragraph()
    _rich(p, "分值锚点：1 明显不足　|　2 有待改进　|　3 基本合格　|　4 良好　|　5 优秀",
          FORM_FONT_PT)

    _para(doc, "", FORM_FONT_PT)
    p = doc.add_paragraph()
    _rich(p, "**教师信息（选填，仅用于区分评分者，不与个人身份关联）**", FORM_FONT_PT)
    t2 = _make_table(doc, 3, 2)
    for ri, (key, label) in enumerate(TEACHER_LABELS.items()):
        _set_cell_lines(t2.rows[ri].cells[0], [label], FORM_FONT_PT, bold=True)
        t2.rows[ri].cells[0].width = Cm(5.4)
        if key == "mixed":
            _put_dropdown(t2.rows[ri].cells[1], "mixed", ["请选择", "是", "否"])
        else:
            _set_cell_lines(t2.rows[ri].cells[1], [PLACEHOLDER_ANSWER], FORM_FONT_PT)


# --------------------------------------------------------------------------
# 评分块
# --------------------------------------------------------------------------
def _score_block(doc, code: str) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(f"★ 请为教案 {code} 打分")
    _font(r, 12.0, True)

    tbl = _make_table(doc, 5, 6)
    for ci, h in enumerate(["维度"] + [f"{c}\n{n}" for c, n, _ in DIMS]):
        _set_cell_lines(tbl.rows[0].cells[ci], h.split("\n"), FORM_FONT_PT, bold=True)

    # 第 1 行：五个下拉打分格
    _set_cell_lines(tbl.rows[1].cells[0], [ROW_LABELS["score"]], FORM_FONT_PT, bold=True)
    for i, dim in enumerate(DIM_CODES, start=1):
        _put_dropdown(tbl.rows[1].cells[i], f"{code}-{dim}", SCORE_OPTIONS)
    _row_height(tbl.rows[1], 420)

    # 第 2 行：Q6 下拉（跨列合并）
    _set_cell_lines(tbl.rows[2].cells[0], [ROW_LABELS["q6"]], FORM_FONT_PT, bold=True)
    q6 = tbl.rows[2].cells[1].merge(tbl.rows[2].cells[5])
    _put_dropdown(q6, f"{code}-Q6", Q6_OPTIONS)
    _row_height(tbl.rows[2], 400)

    # 第 3、4 行：Q7 / Q8 填空
    for ri, key in ((3, "q7"), (4, "q8")):
        _set_cell_lines(tbl.rows[ri].cells[0], [ROW_LABELS[key]], FORM_FONT_PT, bold=True)
        ans = tbl.rows[ri].cells[1].merge(tbl.rows[ri].cells[5])
        _set_cell_lines(ans, [PLACEHOLDER_ANSWER], FORM_FONT_PT)
        _row_height(tbl.rows[ri], 700)

    for ci, w in enumerate((Cm(2.6),) + (Cm(2.7),) * 5):
        for row in tbl.rows:
            try:
                row.cells[ci].width = w
            except IndexError:                             # pragma: no cover
                pass
    _para(doc, "", FORM_FONT_PT)


# --------------------------------------------------------------------------
# 生成
# --------------------------------------------------------------------------
def build_teacher_form(cards: dict, codes: Sequence[str], teacher_label: str,
                       out_path: Path) -> Path:
    """生成一份**自足**的教师评分卡：说明 + 逐份「教案正文 + 评分表」。"""
    doc = Document()
    _init_styles(doc)

    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.0)

    _cover(doc, codes, teacher_label)

    for i, code in enumerate(codes):
        doc.add_page_break()
        _md_into(doc, cards[code])
        _score_block(doc, code)

    doc.add_page_break()
    p = doc.add_paragraph()
    r = p.add_run("填完了，谢谢您！")
    _font(r, 14.0, True)
    for line in [
        "",
        "请保存本文件，改名为「评分卡-您的称呼.docx」，放进**回收**文件夹。",
        "",
        "若您对本次评议有任何疑问或想撤回评分，请联系组织方，无需说明理由。",
    ]:
        _para(doc, line, FORM_FONT_PT)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


# --------------------------------------------------------------------------
# 回收解析
# --------------------------------------------------------------------------
_SCORE_ALIAS = re.compile(r"为教案\s*([A-H])\s*打分")


def _clean_answer(text: str) -> str:
    t = (text or "").replace("\u3000", " ").strip()
    return "" if t in BLANK else t


def _num_or_none(text: str) -> Optional[float]:
    t = _clean_answer(text)
    if not t:
        return None
    m = re.fullmatch(r"(\d+(?:\.\d+)?)", t)
    return float(m.group(1)) if m else None


def extract_form(path: Path) -> dict:
    """读一份填好的评分卡，抽出教师信息 + 每份教案的 D1–D5 / Q6–Q8。"""
    doc = Document(str(path))
    result = {"file": path.name, "teacher": "", "mixed": "",
              "scores": {}, "warnings": []}
    cur: Optional[str] = None

    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            txt = Paragraph(child, doc).text.strip()
            m = _SCORE_ALIAS.search(txt)
            if m:
                cur = m.group(1)
                result["scores"].setdefault(
                    cur, {d: None for d in DIM_CODES} | {"Q6": "", "Q7": "", "Q8": ""})
            continue
        if child.tag != qn("w:tbl"):
            continue
        tbl = Table(child, doc)
        first = _deep_text(tbl.rows[0].cells[0]._tc).strip()

        if first.startswith("教师称呼"):
            for row in tbl.rows:
                lab = _deep_text(row.cells[0]._tc).strip()
                val = _clean_answer(_deep_text(row.cells[1]._tc))
                if lab.startswith("教师称呼"):
                    result["teacher"] = val
                elif lab.startswith("是否带过"):
                    result["mixed"] = _clean_answer(val)
            continue

        if first == "维度" and cur:
            slot = result["scores"][cur]
            for row in tbl.rows:
                lab = _deep_text(row.cells[0]._tc).strip()
                if lab.startswith("得分"):
                    for i, d in enumerate(DIM_CODES, start=1):
                        slot[d] = _num_or_none(_deep_text(row.cells[i]._tc))
                elif lab.startswith("Q6"):
                    slot["Q6"] = _clean_answer(_deep_text(row.cells[1]._tc))
                elif lab.startswith("Q7"):
                    slot["Q7"] = _clean_answer(_deep_text(row.cells[1]._tc))
                elif lab.startswith("Q8"):
                    slot["Q8"] = _clean_answer(_deep_text(row.cells[1]._tc))

    if not result["scores"]:
        result["warnings"].append("没找到任何「为教案 X 打分」区块 —— 是不是发错了文件？")
    for code, slot in result["scores"].items():
        got = [d for d in DIM_CODES if slot[d] is not None]
        if not got:
            result["warnings"].append(f"教案 {code}：D1–D5 一格都没填")
        elif len(got) < len(DIM_CODES):
            miss = [d for d in DIM_CODES if slot[d] is None]
            result["warnings"].append(f"教案 {code}：缺 {'/'.join(miss)}")
        if not slot["Q6"]:
            result["warnings"].append(f"教案 {code}：Q6 未选")
        if len(got) == len(DIM_CODES) and not slot["Q7"]:
            result["warnings"].append(f"教案 {code}：Q7 必答未填")
        for d in DIM_CODES:
            v = slot[d]
            if v is not None and not (1.0 <= v <= 5.0):
                result["warnings"].append(f"教案 {code}：{d} 分值 {v} 越界（应 1–5）")
    return result
